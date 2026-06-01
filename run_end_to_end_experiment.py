import csv
import json
import os
import re
import subprocess
import tempfile
import time
from pathlib import Path

from openai import OpenAI
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
WORKBOOK = ROOT / "tasks_10.xlsx"
OUT_DIR = ROOT / "results_end_to_end"
MODEL = "gpt-4o"
SOLVER = "gecode"
TIMEOUT_SECONDS = 30
MINIZINC_TIME_LIMIT_MS = 10000


SYSTEM_PROMPT = """You generate MiniZinc artifacts directly from a natural language problem description.
Return only valid JSON with exactly two string fields: model_mzn and data_dzn.
Do not explain your answer. Do not include markdown fences.
Generate a complete MiniZinc model.mzn and data.dzn in one end-to-end zero-shot response."""


USER_PROMPT_TEMPLATE = """Create MiniZinc files directly from this natural language problem description.

Use only the description below. Produce:
- model_mzn: complete MiniZinc model text
- data_dzn: complete MiniZinc data text, or an empty string if no separate data file is needed

Problem description:
{problem}
"""


def read_tasks():
    wb = load_workbook(WORKBOOK, data_only=False)
    ws = wb.active
    tasks = []
    for row in range(2, ws.max_row + 1):
        problem = ws.cell(row, 2).value
        if not problem:
            continue
        tasks.append(
            {
                "task_id": f"task_{len(tasks) + 1:03d}",
                "problem": str(problem),
                "reference_model": str(ws.cell(row, 7).value or ""),
                "reference_data": str(ws.cell(row, 8).value or ""),
            }
        )
    return tasks


def minizinc_version():
    proc = subprocess.run(
        ["minizinc", "--version"],
        cwd=ROOT,
        text=True,
        capture_output=True,
        timeout=30,
    )
    return (proc.stdout + proc.stderr).strip()


def generate(client, problem):
    prompt = USER_PROMPT_TEMPLATE.format(problem=problem)
    response = client.chat.completions.create(
        model=MODEL,
        messages=[
            {"role": "system", "content": SYSTEM_PROMPT},
            {"role": "user", "content": prompt},
        ],
        temperature=0,
        response_format={"type": "json_object"},
    )
    content = response.choices[0].message.content or "{}"
    parsed = json.loads(content)
    return {
        "model_mzn": str(parsed.get("model_mzn", "")),
        "data_dzn": str(parsed.get("data_dzn", "")),
        "raw_response": content,
        "prompt": prompt,
    }


def load_or_generate(client, task_dir, problem):
    audit_path = task_dir / "generation_audit.json"
    model_path = task_dir / "generated_model.mzn"
    data_path = task_dir / "generated_data.dzn"
    if audit_path.exists() and model_path.exists() and data_path.exists():
        audit = json.loads(audit_path.read_text(encoding="utf-8"))
        return {
            "model_mzn": model_path.read_text(encoding="utf-8"),
            "data_dzn": data_path.read_text(encoding="utf-8"),
            "raw_response": str(audit.get("raw_response", "")),
            "prompt": str(audit.get("prompt", "")),
        }

    generated = generate(client, problem)
    write_text(model_path, generated["model_mzn"])
    write_text(data_path, generated["data_dzn"])
    write_text(audit_path, json.dumps(generated, indent=2))
    return generated


def run_minizinc(model_path, data_path=None):
    cmd = [
        "minizinc",
        "--solver",
        SOLVER,
        "--statistics",
        "--time-limit",
        str(MINIZINC_TIME_LIMIT_MS),
        str(model_path),
    ]
    if data_path is not None and Path(data_path).read_text(encoding="utf-8").strip():
        cmd.append(str(data_path))
    started = time.perf_counter()
    try:
        proc = subprocess.run(
            cmd,
            cwd=ROOT,
            text=True,
            capture_output=True,
            timeout=TIMEOUT_SECONDS,
        )
        elapsed = time.perf_counter() - started
        output = (proc.stdout or "") + (proc.stderr or "")
        return proc.returncode, output, elapsed, False
    except subprocess.TimeoutExpired as exc:
        elapsed = time.perf_counter() - started
        output = ((exc.stdout or "") + (exc.stderr or ""))
        if isinstance(output, bytes):
            output = output.decode("utf-8", errors="replace")
        return 124, output + f"\nTIMEOUT after {TIMEOUT_SECONDS} seconds\n", elapsed, True


def parse_solve_kind(model_text):
    lowered = model_text.lower()
    if re.search(r"\bsolve\s+(minimize|maximize)\b", lowered):
        return "optimize"
    if re.search(r"\bsolve\s+satisfy\b", lowered):
        return "satisfy"
    return "unknown"


def parse_status(output, returncode, timed_out, solve_kind):
    text = output.upper()
    if timed_out:
        return "TIMEOUT"
    if "UNSATISFIABLE" in text:
        return "UNSATISFIABLE"
    if "UNKNOWN" in text:
        return "UNKNOWN"
    if returncode != 0 or "ERROR:" in text:
        return "ERROR"
    if "==========" in output:
        return "OPTIMAL" if solve_kind == "optimize" else "SATISFIABLE"
    if "----------" in output:
        return "SATISFIABLE"
    return "SUCCESS"


def comparable_status(status):
    if status in {"OPTIMAL", "SATISFIABLE", "SUCCESS"}:
        return "SAT"
    if status == "UNSATISFIABLE":
        return "UNSAT"
    return status


def parse_objective(output):
    matches = re.findall(r"%%%mzn-stat:\s*objective=([^\s]+)", output)
    return matches[-1] if matches else ""


def strip_stats(output):
    lines = []
    for line in output.splitlines():
        if line.startswith("%%%mzn-stat:"):
            continue
        if line.startswith("%%%mzn-stat-end"):
            continue
        lines.append(line.strip())
    return "\n".join(line for line in lines if line)


def compare_results(g_status, r_status, g_output, r_output):
    g_comp = comparable_status(g_status)
    r_comp = comparable_status(r_status)
    if g_comp in {"ERROR", "TIMEOUT", "UNKNOWN"} or r_comp in {"ERROR", "TIMEOUT", "UNKNOWN"}:
        return "NOT_COMPARABLE"
    if g_comp != r_comp:
        return "DIFFERENT_RESULT"
    if g_comp == "UNSAT":
        return "SAME_RESULT"

    g_obj = parse_objective(g_output)
    r_obj = parse_objective(r_output)
    if g_obj and r_obj:
        return "SAME_RESULT" if g_obj == r_obj else "DIFFERENT_RESULT"

    g_solution = strip_stats(g_output)
    r_solution = strip_stats(r_output)
    if g_solution and r_solution and g_solution == r_solution:
        return "SAME_RESULT"
    return "NOT_COMPARABLE"


def write_text(path, text):
    Path(path).write_text(text or "", encoding="utf-8", newline="\n")


def main():
    if not os.getenv("OPENAI_API_KEY"):
        raise RuntimeError("OPENAI_API_KEY is not set")

    version = minizinc_version()
    tasks = read_tasks()
    OUT_DIR.mkdir(exist_ok=True)

    client = OpenAI()
    rows = []

    for task in tasks:
        task_dir = OUT_DIR / task["task_id"]
        task_dir.mkdir(exist_ok=True)
        write_text(task_dir / "problem.txt", task["problem"])

        generated = load_or_generate(client, task_dir, task["problem"])

        g_return, g_output, g_elapsed, g_timeout = run_minizinc(
            task_dir / "generated_model.mzn",
            task_dir / "generated_data.dzn",
        )
        g_status = parse_status(
            g_output,
            g_return,
            g_timeout,
            parse_solve_kind(generated["model_mzn"]),
        )
        write_text(task_dir / "generated_output.txt", g_output)

        with tempfile.TemporaryDirectory(dir=OUT_DIR) as tmp:
            tmp_path = Path(tmp)
            ref_model = tmp_path / "reference_model.mzn"
            ref_data = tmp_path / "reference_data.dzn"
            write_text(ref_model, task["reference_model"])
            write_text(ref_data, task["reference_data"])
            r_return, r_output, _, r_timeout = run_minizinc(ref_model, ref_data)

        r_status = parse_status(
            r_output,
            r_return,
            r_timeout,
            parse_solve_kind(task["reference_model"]),
        )
        write_text(task_dir / "reference_output.txt", r_output)

        valid_model = g_status not in {"ERROR", "TIMEOUT", "UNKNOWN"}
        comparison = compare_results(g_status, r_status, g_output, r_output)
        notes = []
        if not valid_model:
            notes.append("Generated model did not compile/execute successfully.")
        if r_status in {"ERROR", "TIMEOUT", "UNKNOWN"}:
            notes.append("Reference run did not produce a comparable result.")
        if comparison == "NOT_COMPARABLE" and valid_model:
            notes.append("SAT/UNSAT matched or was unavailable, but objective/solution comparison was not meaningful.")

        rows.append(
            {
                "task_id": task["task_id"],
                "valid_model": "TRUE" if valid_model else "FALSE",
                "generated_solver_status": g_status,
                "reference_solver_status": r_status,
                "comparison_result": comparison,
                "execution_time_seconds": f"{g_elapsed:.3f}",
                "notes": " ".join(notes),
            }
        )

    final_csv = OUT_DIR / "final_results.csv"
    with final_csv.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)

    total = len(rows)
    valid = sum(1 for row in rows if row["valid_model"] == "TRUE")
    same = sum(1 for row in rows if row["comparison_result"] == "SAME_RESULT")
    avg_time = sum(float(row["execution_time_seconds"]) for row in rows) / total if total else 0.0
    summary = {
        "total_tasks": total,
        "valid_models": valid,
        "valid_percentage": f"{(valid / total * 100) if total else 0:.2f}",
        "same_results": same,
        "same_results_percentage": f"{(same / total * 100) if total else 0:.2f}",
        "average_execution_time_seconds": f"{avg_time:.3f}",
    }
    with (OUT_DIR / "summary_statistics.csv").open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(summary.keys()))
        writer.writeheader()
        writer.writerow(summary)

    report_lines = [
        "# End-to-End Zero-Shot MiniZinc Generation Baseline",
        "",
        "## Methodology",
        "",
        "For each task, only column B from `tasks_10.xlsx` was provided to GPT-4o. The model was asked once to directly generate `model.mzn` and `data.dzn` as JSON fields. No intermediate representations, semantic extraction, decomposition, repair loop, recognizer, parser, or iterative workflow was used. Columns G and H were used only after generation for reference evaluation.",
        "",
        f"Generated models were executed with `minizinc --solver gecode --statistics --time-limit {MINIZINC_TIME_LIMIT_MS} generated_model.mzn generated_data.dzn`. Reference models were executed with the same solver, statistics flag, and time limit.",
        "",
        "## Assumptions",
        "",
        "- A generated model is counted as valid when MiniZinc compiles and executes without timeout, unknown status, or error.",
        "- SAT/UNSAT status is compared first. When both runs are satisfiable and MiniZinc reports objective values, objectives are compared. If objectives are absent and solution text is not directly comparable, the result is marked `NOT_COMPARABLE`.",
        f"- MiniZinc solver time limit per run: {MINIZINC_TIME_LIMIT_MS / 1000:.0f} seconds.",
        f"- Outer process timeout per run: {TIMEOUT_SECONDS} seconds.",
        "",
        "## Configuration",
        "",
        f"- LLM: `{MODEL}`",
        "- Temperature: `0`",
        "- Interaction pattern: one API call per task",
        "- Solver: `gecode`",
        f"- MiniZinc version: `{version.splitlines()[0] if version else 'unknown'}`",
        "",
        "## Per-Task Outcomes",
        "",
        "| task_id | valid_model | generated_status | reference_status | comparison | time_s | notes |",
        "|---|---:|---|---|---|---:|---|",
    ]
    for row in rows:
        report_lines.append(
            f"| {row['task_id']} | {row['valid_model']} | {row['generated_solver_status']} | "
            f"{row['reference_solver_status']} | {row['comparison_result']} | "
            f"{row['execution_time_seconds']} | {row['notes']} |"
        )
    report_lines.extend(
        [
            "",
            "## Aggregate Statistics",
            "",
            f"- Total tasks: {summary['total_tasks']}",
            f"- Valid models: {summary['valid_models']} ({summary['valid_percentage']}%)",
            f"- Same results: {summary['same_results']} ({summary['same_results_percentage']}%)",
            f"- Average generated-model execution time: {summary['average_execution_time_seconds']} seconds",
            "",
            "## Failure Analysis",
            "",
        ]
    )
    invalid_rows = [row for row in rows if row["valid_model"] == "FALSE"]
    diff_rows = [row for row in rows if row["comparison_result"] == "DIFFERENT_RESULT"]
    nc_rows = [row for row in rows if row["comparison_result"] == "NOT_COMPARABLE"]
    if not invalid_rows and not diff_rows and not nc_rows:
        report_lines.append("No failures or non-comparable outcomes were observed.")
    else:
        report_lines.append(f"- Invalid generated models: {len(invalid_rows)}")
        report_lines.append(f"- Different results: {len(diff_rows)}")
        report_lines.append(f"- Not comparable: {len(nc_rows)}")
        for row in rows:
            if row["valid_model"] == "FALSE" or row["comparison_result"] != "SAME_RESULT":
                report_lines.append(
                    f"- {row['task_id']}: valid={row['valid_model']}, "
                    f"generated={row['generated_solver_status']}, "
                    f"reference={row['reference_solver_status']}, "
                    f"comparison={row['comparison_result']}. {row['notes']}"
                )

    write_text(OUT_DIR / "experiment_report.md", "\n".join(report_lines) + "\n")
    print(f"Wrote results to {OUT_DIR}")


if __name__ == "__main__":
    main()
